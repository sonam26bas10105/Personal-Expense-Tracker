"""Data-access layer: the only module that talks to the Excel workbook.

Other modules use these helpers, so the storage engine (Excel today, possibly
SQLite later) can be replaced by changing this file alone.
"""
import os

from openpyxl import Workbook, load_workbook

from logger import get_logger

logger = get_logger(__name__)

DB_NAME = os.path.join(os.path.dirname(os.path.abspath(__file__)), "expenses.xlsx")

SHEETS = {
    "Expenses": ["ID", "Amount", "Category", "Description", "Date"],
    "Categories": ["ID", "Name"],
    "Budget": ["Amount"],
}


class StorageError(Exception):
    """Raised when the workbook cannot be created, opened or saved."""


def create_tables(path: str = None) -> None:
    """Create the workbook and any missing sheets (with header rows)."""
    path = path or DB_NAME
    try:
        if os.path.exists(path):
            wb = load_workbook(path)
        else:
            wb = Workbook()
            wb.remove(wb.active)

        changed = False
        for name, headers in SHEETS.items():
            if name not in wb.sheetnames:
                wb.create_sheet(name).append(headers)
                changed = True

        if changed:
            wb.save(path)
    except OSError as exc:
        logger.error("Failed to create workbook %s: %s", path, exc)
        raise StorageError(str(exc)) from exc

    logger.info("Workbook verified/created at %s", path)


def connect(path: str = None) -> Workbook:
    """Open the workbook, creating it first if it does not exist."""
    path = path or DB_NAME
    try:
        if not os.path.exists(path):
            create_tables(path)
        return load_workbook(path)
    except OSError as exc:
        logger.error("Failed to open workbook %s: %s", path, exc)
        raise StorageError(str(exc)) from exc


def save_book(wb: Workbook, path: str = None) -> None:
    """Save the workbook; a locked file becomes a friendly StorageError."""
    path = path or DB_NAME
    try:
        wb.save(path)
    except OSError as exc:
        logger.error("Failed to save workbook %s: %s", path, exc)
        raise StorageError(
            "could not write to the Excel file. Close it if it is open in Excel."
        ) from exc


def data_rows(ws) -> list:
    """Return every data row (header skipped, blank rows ignored) as tuples."""
    return [tuple(row) for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None]


def next_id(ws) -> int:
    """Next free id: highest existing id + 1 (stays unique after deletions)."""
    return max((row[0] for row in data_rows(ws)), default=0) + 1


def find_row(ws, record_id: int):
    """Return the worksheet row number holding record_id, or None."""
    for row in ws.iter_rows(min_row=2):
        if row[0].value == record_id:
            return row[0].row
    return None
